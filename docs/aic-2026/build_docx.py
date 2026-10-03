#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把 content.md 排版为符合《2026AIC-算法创新赛-作品方案模板》格式规范的技术报告。

格式规范（摘自模板）：
  字体:宋体
  目录标题 二号(22pt) 粗体 / 一级标题 三号(16pt) 粗体 / 二级标题 四号(14pt) 粗体
  三级标题 小四(12pt) 粗体 / 正文 小四(12pt)
  行距:单倍行距
  页边距:上2.5cm 下2.5cm 左3cm 右3cm 装订线0 / 页眉1.5cm 页脚1.5cm
  纸型:A4 纵向

用法：python3 build_docx.py
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / "content.md"
OUT = HERE / "银龄智办-技术报告-2026AIC.docx"

FONT = "宋体"
MONO = "Consolas"

# 中文字号 -> pt
SIZE = {
    "初号": 42, "小初": 36, "一号": 26, "小一": 24, "二号": 22, "小二": 18,
    "三号": 16, "小三": 15, "四号": 14, "小四": 12, "五号": 10.5, "小五": 9,
}

H1, H2, H3 = SIZE["三号"], SIZE["四号"], SIZE["小四"]
BODY = SIZE["小四"]
TABLE_TXT = SIZE["五号"]
NOTE_TXT = SIZE["小五"]


# ---------------------------------------------------------------- 基础工具

def set_font(run, name=FONT, size=None, bold=None, italic=None, color=None):
    """同时设置 ascii / hAnsi / eastAsia 字体，中文才不会回退成 Times。"""
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def para_spacing(p, before=0, after=0, line=1.0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE if line == 1.0 else WD_LINE_SPACING.MULTIPLE
    return p


TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+?`)")


def add_runs(p, text, size=BODY, bold=False, name=FONT):
    """解析 **粗体** 与 `等宽` 行内标记。"""
    for part in TOKEN_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = p.add_run(part[2:-2])
            set_font(r, name, size, True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = p.add_run(part[1:-1])
            set_font(r, MONO, size * 0.94, bold)
        else:
            r = p.add_run(part)
            set_font(r, name, size, bold)
    return p


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    set_font(run, FONT, BODY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "（在 Word / WPS 中按 F9 或右键“更新域”生成）"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, placeholder, end):
        run._r.append(node)


# ---------------------------------------------------------------- 文档骨架

def build_document():
    doc = Document()

    # 页面设置：A4 纵向 + 模板页边距
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(3)
    sec.header_distance = Cm(1.5)
    sec.footer_distance = Cm(1.5)

    # Normal：正文小四、单倍行距、宋体
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    # 标题样式：宋体 + 模板字号 + 粗体 + 黑色（Heading 样式自带大纲级别，目录域才认得）
    for name, size in (("Heading 1", H1), ("Heading 2", H2), ("Heading 3", H3)):
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 6)
        st.paragraph_format.space_after = Pt(6 if name == "Heading 1" else 4)
        st.paragraph_format.keep_with_next = True

    # 页眉
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(hp, "2026 第八届全球校园人工智能算法精英大赛", NOTE_TXT)

    # 页脚：页码域
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE   \\* MERGEFORMAT")

    # 打开文档时自动更新域（目录页码无需手工 F9）
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)

    return doc


def build_cover(doc, lines):
    """封面：大赛名 / 赛道 / 赛题 / 报告名 / 信息栏。"""
    def blank(times=1, size=BODY):
        for _ in range(times):
            p = doc.add_paragraph()
            para_spacing(p, 0, 0)
            r = p.add_run("")
            set_font(r, FONT, size)

    blank(2)
    for text in lines["head"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para_spacing(p, 0, 6)
        add_runs(p, text, SIZE["二号"], True)

    if lines["track"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para_spacing(p, 6, 0)
        add_runs(p, lines["track"], SIZE["三号"], True)

    blank(2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para_spacing(p, 0, 0)
    add_runs(p, lines["doctitle"], SIZE["一号"], True)
    blank(2)

    for info in lines["info"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para_spacing(p, 0, 10)
        add_runs(p, info, SIZE["四号"])


# ---------------------------------------------------------------- 正文渲染

def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True

    for i, text in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para_spacing(p, 2, 2)
        add_runs(p, text, TABLE_TXT, True)
        shade(cell, "F2F2F2")

    for row in body:
        cells = t.add_row().cells
        for i, text in enumerate(row[: len(header)]):
            cell = cells[i]
            cell.text = ""
            first = True
            for seg in text.split("<br>"):
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                para_spacing(p, 1, 1)
                add_runs(p, seg, TABLE_TXT)

    # 表后留一点空隙
    spacer = doc.add_paragraph()
    para_spacing(spacer, 0, 0)
    set_font(spacer.add_run(""), FONT, SIZE["五号"])
    return t


def render(doc, text):
    lines = text.split("\n")
    i = 0
    pending_table = []

    def flush_table():
        nonlocal pending_table
        if pending_table:
            add_table(doc, pending_table)
            pending_table = []

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        # 表格行
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                pending_table.append(cells)
            i += 1
            continue
        flush_table()

        if not stripped:
            i += 1
            continue

        if stripped == "[END]":
            break
        if stripped == "[PAGEBREAK]":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            i += 1
            continue
        if stripped == "[TOC]":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para_spacing(p, 0, 12)
            add_runs(p, "目　录", SIZE["二号"], True)
            p2 = doc.add_paragraph()
            para_spacing(p2, 0, 0)
            add_field(p2, 'TOC \\o "1-3" \\h \\z \\u')
            i += 1
            continue

        # 代码块
        if stripped.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                p = doc.add_paragraph()
                para_spacing(p, 0, 0)
                p.paragraph_format.left_indent = Cm(0.6)
                r = p.add_run(lines[i])
                set_font(r, MONO, NOTE_TXT)
                i += 1
            i += 1
            continue

        # 图位标注
        if stripped.startswith("[FIG]"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para_spacing(p, 6, 10)
            add_runs(p, stripped[5:].strip(), NOTE_TXT, False)
            i += 1
            continue

        # 标题
        if stripped.startswith("### "):
            p = doc.add_heading(level=3)
            add_runs(p, stripped[4:], H3, True)
            para_spacing(p, 6, 4)
            i += 1
            continue
        if stripped.startswith("## "):
            p = doc.add_heading(level=2)
            add_runs(p, stripped[3:], H2, True)
            para_spacing(p, 8, 4)
            i += 1
            continue
        if stripped.startswith("# "):
            p = doc.add_heading(level=1)
            add_runs(p, stripped[2:], H1, True)
            para_spacing(p, 12, 6)
            i += 1
            continue

        # 无序列表
        if stripped.startswith("- "):
            p = doc.add_paragraph()
            para_spacing(p, 0, 2)
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.first_line_indent = Cm(-0.37)
            add_runs(p, "· " + stripped[2:], BODY)
            i += 1
            continue

        # 普通正文（首行缩进 2 字符）
        p = doc.add_paragraph()
        para_spacing(p, 0, 3)
        p.paragraph_format.first_line_indent = Pt(BODY * 2)
        add_runs(p, stripped, BODY)
        i += 1

    flush_table()


def main():
    src = SRC.read_text(encoding="utf-8")

    # 拆出封面块
    cover_lines = {"head": [], "track": "", "doctitle": "", "info": []}
    body_lines = []
    in_cover = False
    for line in src.split("\n"):
        s = line.strip()
        if s == "[COVER]":
            in_cover = True
            continue
        if in_cover:
            if s == "[PAGEBREAK]":
                in_cover = False
                body_lines.append("[PAGEBREAK]")
                continue
            if s.startswith("[TRACK]"):
                cover_lines["track"] = s[len("[TRACK]"):].strip()
            elif s.startswith("[DOCTITLE]"):
                cover_lines["doctitle"] = s[len("[DOCTITLE]"):].strip()
            elif s.startswith("[INFO]"):
                cover_lines["info"].append(s[len("[INFO]"):].strip())
            elif s:
                cover_lines["head"].append(s)
            continue
        body_lines.append(line)

    doc = build_document()
    build_cover(doc, cover_lines)
    render(doc, "\n".join(body_lines))
    doc.save(OUT)

    # ---- 自检 ----
    m = re.search(r"# 作品简介\n\n(.+?)\n\n#", src, re.S)
    intro = re.sub(r"\s", "", m.group(1)) if m else ""
    print(f"输出文件 : {OUT}")
    print(f"作品简介 : {len(intro)} 字（模板要求 ≤300）")
    print(f"总字符数 : {len(re.sub(chr(92) + 's', '', src))}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
