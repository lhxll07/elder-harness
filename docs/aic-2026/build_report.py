#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把《作品方案-定稿.md》排版成符合 2026AIC 作品方案模板格式规范的 Word 文档。

模板格式规范：
  字体:宋体 | 目录标题 二号粗体 | 一级标题 三号粗体 | 二级标题 四号粗体
  三级标题 小四粗体 | 正文 小四 | 单倍行距
  页边距 上下2.5cm 左右3cm | 页眉页脚1.5cm | A4 纵向

图表规范：表题在表上方、图题在图下方，全文按出现顺序统一编号。
Mermaid 图需先渲染为 figures/figN.png（见 render_figures.sh），本脚本只负责插入。

用法：python3 build_report.py
"""

import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / "作品方案-定稿.md"
OUT = HERE / "银龄智办-作品方案-2026AIC-AI+软件创新.docx"
FIGDIR = HERE / "figures"

FONT, MONO = "宋体", "Consolas"
# 中西文混排：中文用 FONT（eastAsia），西文/数字/半角标点用 Times New Roman（ascii/hAnsi）。
# 这是 Word 的标准做法——靠 rFonts 分槽，不必把一行文字拆成多个 run。
LATIN = "Times New Roman"
# 封面照搬大赛模板：标题区黑体、信息区仿宋；Logo 取自模板页眉
HEI, FANG = "黑体", "仿宋"
LOGO = HERE / "assets" / "aic-logo-header.png"   # 已按模板 srcRect 裁掉底部年份文字
HEADER_TEXT = "2026第八届全球校园人工智能算法精英大赛"
S = {"一号": 26, "二号": 22, "三号": 16, "四号": 14, "小四": 12, "五号": 10.5, "小五": 9}
H1, H2, H3, BODY = S["三号"], S["四号"], S["小四"], S["小四"]
TBL, NOTE = S["五号"], S["小五"]

# ---------------------------------------------------------------- 图表题注
# 按文档中出现顺序一一对应；数量不符时脚本报错，避免编号错位。

TABLE_CAPTIONS = [
    "作品概览一览",
    "核心证据索引（给评审的定位表）",
    "老年人手机办事痛点的四层十二项剖析",
    "现有助老方案对比（一）：执行能力维度",
    "现有助老方案对比（二）：工程与合规维度",
    "面向老年人的需求清单与功能对应",
    "核心功能与实现状态",
    "研发目标与当前达成情况",
    "系统分层模块职责与数据交互",
    "关键技术选型与依据",
    "四条技术路线的模型选型对比",
    "按可逆性划分的三级执行策略",
    "已有工作与本作品的技术差异",
    "三条机制在相邻领域的先行工作与边界",
    "八项创新点及其差异与支撑证据",
    "智能体可调用的工具目录",
    "五种任务收尾方式及其对老人的呈现",
    "典型任务全流程走查（美团点餐至结算交接）",
    "跨应用悬浮接线台的交互要素与工程细节",
    "多应用适配路径与真机任务覆盖",
    "真机问题驱动的开发迭代阶段",
    "代码规模统计",
    "测试分层与覆盖内容",
    "真机任务测试记录（小样本）",
    "测试发现问题与改进效果",
    "已修复的安全与可靠性问题",
    "待改进问题与计划",
    "运行环境要求",
    "权限与系统适配说明",
    "论文问题的三条判定标准",
    "实验一（说服不变性）成对诊断结果",
    "实验二（能力饱和对照）结果",
    "实验三（声明—动作缺口与跨轮漂移）结果",
    "实验四（环境注入）修复前的两个指标",
    "实验四修复后的结果与反向用例",
    "实验五（隐蔽型危险动作）的词表覆盖与边界",
    "执行约束的三组对照",
    "F1—F4 四项修复的前后对照",
    "同一目标在三个电商 App 上的结果",
    "模糊指令与明确指令下的澄清行为",
    "火车票四项答案与页面逐项核对",
    "坐标定位命中率：三个屏幕",
    "坐标定位命中率的空间分布",
    "真机实验发现的三个新问题",
    "典型应用场景与交接边界",
    "插图清单与作图说明",
    "关键设计参数",
    "上下文前缀缓存命中实测",
    "文献验证清单与本报告编号的对应关系",
    "第三方组件与服务清单",
]

FIGURE_CAPTIONS = [
    "痛点四层因果闭环与四项机制的补位位置",
    "系统总体架构（四层）",
    "完成声明核验的三值判定流程",
    "一个判据的三层复用：屏幕动作 / 业务动作 / 导航动作",
    "功能架构：发起 / 办事 / 保障 / 接力",
    "任务执行流程与五种收尾方式",
    "观察层的三条路径与自动切换",
    "平安守望的判定状态机（含“绝不狼来了”的四道闸门）",
]

# ---------------------------------------------------------------- 参考文献
# 正文用 [@key] 引用，编号由本表顺序自动决定（分组编号：政策 / 学术 / 技术文档）。
# 增删文献只改这里，正文编号自动跟随，不会出现"改了文献表正文编号全错"的问题。
# 所有条目均经逐条抓取核对（见 refs/ 目录下的验证清单）。

REFS = [
    # —— 一、政策与统计文献 ——
    ("p_stats2024", "国家统计局.《2024 年国民经济和社会发展统计公报》及 2025 年 1 月人口数据解读（王萍萍）. 2025-01-17. https://www.stats.gov.cn/sj/sjjd/202501/t20250117_1958337.html"),
    ("p_stats2025", "国家统计局. 2025 年全国人口数据解读（全国 1% 人口抽样调查）. 2026-01-19. https://www.stats.gov.cn/zt_18555/zthd/lhfw/2026lhzt/2026sjjd/202602/t20260202_1962437.html"),
    ("p_cnnic55", "中国互联网络信息中心（CNNIC）. 第 55 次《中国互联网络发展状况统计报告》. 2025-01-17. https://www.cnnic.cn/n4/2025/0117/c88-11229.html"),
    ("p_cnnic56", "中国互联网络信息中心（CNNIC）. 第 56 次《中国互联网络发展状况统计报告》（发布稿）. 2025-07-21. https://www.cnnic.net.cn/n4/2025/0721/c326-11327.html ；报告全文 PDF：https://www.cnnic.net.cn/NMediaFile/2025/0730/MAIN1753846666507QEK67ZS9DH.pdf"),
    ("p_aging2024", "民政部、全国老龄办.《2024 年度国家老龄事业发展公报》. 2025. https://www.mca.gov.cn/n152/n165/c1662004999980006089/part/21508.pdf"),
    ("p_survey5", "民政部、全国老龄办等.《第五次中国城乡老年人生活状况抽样调查基本数据公报》. 2024-10-17（标准时点 2021-08-01）. http://www.crca.cn/index.php/19-data-resource/life/1117-2024-10-17-08-01-05.html"),
    ("p_class2024", "Hu, H., Xu, W. Socioeconomic differences in digital inequality among Chinese older adults: Results from a nationally representative sample. PLoS ONE 19(4): e0300433, 2024. https://doi.org/10.1371/journal.pone.0300433"),
    ("p_a11ylaw", "全国人民代表大会常务委员会.《中华人民共和国无障碍环境建设法》（2023-06-28 通过，2023-09-01 施行）. https://www.gov.cn/yaowen/liebiao/202306/content_6888910.htm"),
    ("p_gwy45", "国务院办公厅.《关于切实解决老年人运用智能技术困难的实施方案》（国办发〔2020〕45 号）. 2020-11-24. https://www.gov.cn/zhengce/content/2020-11/24/content_5563804.htm"),
    ("p_gwy1", "国务院办公厅.《关于发展银发经济增进老年人福祉的意见》（国办发〔2024〕1 号）. 2024-01-11. https://www.gov.cn/zhengce/content/202401/content_6926087.htm"),
    ("p_gwy35", "国务院.《“十四五”国家老龄事业发展和养老服务体系规划》（国发〔2021〕35 号）. 2022-02-21. https://www.gov.cn/zhengce/content/2022-02/21/content_5674844.htm"),
    ("p_miit200", "工业和信息化部.《互联网应用适老化及无障碍改造专项行动方案》（工信部信管〔2020〕200 号）. 2020-12-24. https://www.gov.cn/zhengce/zhengceku/2020-12/26/content_5573472.htm"),
    ("p_miit251", "工业和信息化部.《促进数字技术适老化高质量发展工作方案》（工信部信管〔2023〕251 号）. 2023-12-19. https://www.gov.cn/zhengce/zhengceku/202312/content_6922847.htm"),
    ("p_smarthealth", "工业和信息化部、民政部、国家卫生健康委.《智慧健康养老产品及服务推广目录（2024 年版）》（工信部联电子函〔2025〕129 号）. 2025-06-05. https://www.miit.gov.cn/zwgk/zcwj/wjfb/tg/art/2025/art_bfa9511fecfb4b22b804f82d5f864690.html"),
    ("p_15th5", "《中华人民共和国国民经济和社会发展第十五个五年规划纲要》第四十章“积极应对人口老龄化”. 2026-03-12 批准，新华社 2026-03-13 受权全文播发."),
    ("p_silver", "中国老龄科学研究中心测算；参见《银发经济蓝皮书：中国银发经济发展报告（2024）》及人民网、北京日报报道（2024-12）：我国银发经济规模约 7 万亿元、占 GDP 约 6%，预计 2035 年达 30 万亿元、占 GDP 约 10%. https://m.gmw.cn/2024-12/24/content_1303930531.htm"),
    ("p_jjrb", "经济日报.《从养老到享老：重塑银发经济的价值坐标》. 2026-07. http://www.jingjiribao.cn/static/detail.jsp?id=679506"),
    ("p_antifraud", "民政部网站转载《中国老年报》.《三种涉老诈骗高发 老年人如何防范》（通报公安部新闻发布会）. 2026-06-16. https://www.mca.gov.cn/n1288/n1290/n1316/c1662004999980011478/content.html"),

    # —— 二、学术文献：手机端 GUI 智能体 ——
    ("a_mobileagent", "Wang, J., Xu, H., Ye, J., et al. Mobile-Agent: Autonomous Multi-Modal Mobile Device Agent with Visual Perception. 2024. arXiv:2401.16158. https://arxiv.org/abs/2401.16158"),
    ("a_mobileagentv2", "Wang, J., Xu, H., Jia, H., et al. Mobile-Agent-v2: Mobile Device Operation Assistant with Effective Navigation via Multi-Agent Collaboration. 2024. arXiv:2406.01014. https://arxiv.org/abs/2406.01014"),
    ("a_mobileagentv3", "Ye, J., Zhang, X., Xu, H., et al. Mobile-Agent-v3: Fundamental Agents for GUI Automation. 2025. arXiv:2508.15144. https://arxiv.org/abs/2508.15144"),
    ("a_qwenuiagent", "Zhou, H., Tong, P., Zhang, X., et al. Qwen-UI-Agent Technical Report: Toward Next-Generation Real-World Centric Foundation GUI Agents. 2026. arXiv:2607.28227. https://arxiv.org/abs/2607.28227"),
    ("a_memgui", "Liu, G., Wu, G., Liu, C., et al. MemGUI-Agent: An End-to-End Long-Horizon Mobile GUI Agent with Proactive Context Management. 2026. arXiv:2606.19926. https://arxiv.org/abs/2606.19926"),
    ("a_sametasks", "Tran, T., Koh, N., Matsunaga, D. E., et al. Same Tasks, Different Apps: Why Mobile GUI Agents Fail to Generalize? Findings of EMNLP 2026. arXiv:2609.34139. https://arxiv.org/abs/2609.34139"),
    ("a_appagent", "Zhang, C., Yang, Z., Liu, J., et al. AppAgent: Multimodal Agents as Smartphone Users. CHI 2025. arXiv:2312.13771. https://arxiv.org/abs/2312.13771"),
    ("a_autodroid", "Wen, H., Li, Y., Liu, G., et al. AutoDroid: LLM-powered Task Automation in Android. MobiCom 2024. arXiv:2308.15272. https://arxiv.org/abs/2308.15272"),
    ("a_androidworld", "Rawles, C., Clinckemaillie, S., Chang, Y., et al. AndroidWorld: A Dynamic Benchmarking Environment for Autonomous Agents. 2024. arXiv:2405.14573. https://arxiv.org/abs/2405.14573"),
    ("a_mobileagentbench", "Wang, L., Deng, Y., Zha, Y., et al. MobileAgentBench: An Efficient and User-Friendly Benchmark for Mobile LLM Agents. 2024. arXiv:2406.08184. https://arxiv.org/abs/2406.08184"),

    # —— 二、学术文献：网页与通用 GUI 智能体 ——
    ("a_mind2web", "Deng, X., Gu, Y., Zheng, B., et al. Mind2Web: Towards a Generalist Agent for the Web. NeurIPS 2023 Spotlight. arXiv:2306.06070. https://arxiv.org/abs/2306.06070"),
    ("a_webarena", "Zhou, S., Xu, F. F., Zhu, H., et al. WebArena: A Realistic Web Environment for Building Autonomous Agents. 2023. arXiv:2307.13854. https://arxiv.org/abs/2307.13854"),
    ("a_osworld", "Xie, T., Zhang, D., Chen, J., et al. OSWorld: Benchmarking Multimodal Agents for Open-Ended Tasks in Real Computer Environments. 2024. arXiv:2404.07972. https://arxiv.org/abs/2404.07972"),
    ("a_seeact", "Zheng, B., Gou, B., Kil, J., et al. GPT-4V(ision) is a Generalist Web Agent, if Grounded. 2024. arXiv:2401.01614. https://arxiv.org/abs/2401.01614"),
    ("a_webvoyager", "He, H., Yao, W., Ma, K., et al. WebVoyager: Building an End-to-End Web Agent with Large Multimodal Models. ACL 2024. arXiv:2401.13919. https://arxiv.org/abs/2401.13919"),
    ("a_websuite", "Li, E., Waldo, J. WebSuite: Systematically Evaluating Why Web Agents Fail. 2024. arXiv:2406.01623. https://arxiv.org/abs/2406.01623"),

    # —— 二、学术文献：GUI 元素定位（Grounding） ——
    ("a_setofmark", "Yang, J., Zhang, H., Li, F., et al. Set-of-Mark Prompting Unleashes Extraordinary Visual Grounding in GPT-4V. 2023. arXiv:2310.11441. https://arxiv.org/abs/2310.11441"),
    ("a_ferretui", "You, K., Zhang, H., Schoop, E., et al. Ferret-UI: Grounded Mobile UI Understanding with Multimodal LLMs. 2024. arXiv:2404.05719. https://arxiv.org/abs/2404.05719"),
    ("a_ferretui2", "Li, Z., You, K., Zhang, H., et al. Ferret-UI 2: Mastering Universal User Interface Understanding Across Platforms. ICLR 2025. arXiv:2410.18967. https://arxiv.org/abs/2410.18967"),
    ("a_seeclick", "Cheng, K., Sun, Q., Chu, Y., et al. SeeClick: Harnessing GUI Grounding for Advanced Visual GUI Agents. 2024. arXiv:2401.10935. https://arxiv.org/abs/2401.10935"),
    ("a_uground", "Gou, B., Wang, R., Zheng, B., et al. Navigating the Digital World as Humans Do: Universal Visual Grounding for GUI Agents. ICLR 2025 Oral. arXiv:2410.05243. https://arxiv.org/abs/2410.05243"),
    ("a_osatlas", "Wu, Z., Wu, Z., Xu, F., et al. OS-ATLAS: A Foundation Action Model for Generalist GUI Agents. 2024. arXiv:2410.23218. https://arxiv.org/abs/2410.23218"),
    ("a_uitars", "Qin, Y., Ye, Y., Fang, J., et al. UI-TARS: Pioneering Automated GUI Interaction with Native Agents. 2025. arXiv:2501.12326. https://arxiv.org/abs/2501.12326"),
    ("a_screenhaystack", "Li, C., Sun, X., Deng, Y., et al. ScreenHaystack: Finding Blind Zones in GUI Grounding. EMNLP 2026 Main Conference. arXiv:2609.32036. https://arxiv.org/abs/2609.32036"),
    ("a_halldfree", "Li, Y., Hou, X. Hallucination-Free GUI Grounding via Regression-Free Layout-Aware Matching. 2026. arXiv:2608.09654. https://arxiv.org/abs/2608.09654"),
    ("a_uq_cua", "Kumar, D., Tayebati, S., Naik, D., et al. Uncertainty Quantification for Computer-Use Agents: A Benchmark across Vision-Language Models and GUI Grounding Datasets. NeurIPS 2026. arXiv:2606.25760. https://arxiv.org/abs/2606.25760"),

    # —— 二、学术文献：移动端智能体评测基准（2025—2026） ——
    ("a_mobileworldsafety", "Chen, S., Li, L., Du, T., Shao, J. MobileWorldSafety: Benchmarking GUI Agent Safety Against Environmental Injection Attacks in Android Apps. 2026. arXiv:2608.17659. https://arxiv.org/abs/2608.17659"),
    ("a_gui_ceval", "Li, Y., Liu, Y., Lu, H., et al. GUI-CEval: A Hierarchical and Comprehensive Chinese Benchmark for Mobile GUI Agents. CVPR 2026. arXiv:2603.15039. https://arxiv.org/abs/2603.15039"),

    # —— 二、学术文献：智能体可靠性、幻觉与安全 ——
    ("a_toolemu", "Ruan, Y., Dong, H., Wang, A., et al. Identifying the Risks of LM Agents with an LM-Emulated Sandbox. 2023. arXiv:2309.15817. https://arxiv.org/abs/2309.15817"),
    ("a_selfcorrect", "Huang, J., Chen, X., Mishra, S., et al. Large Language Models Cannot Self-Correct Reasoning Yet. ICLR 2024. arXiv:2310.01798. https://arxiv.org/abs/2310.01798"),
    ("a_hallucination", "Ye, H., Liu, T., Zhang, A., et al. Cognitive Mirage: A Review of Hallucinations in Large Language Models. 2023. arXiv:2309.06794. https://arxiv.org/abs/2309.06794"),
    ("a_asb", "Zhang, H., Huang, J., Mei, K., et al. Agent Security Bench (ASB): Formalizing and Benchmarking Attacks and Defenses in LLM-based Agents. ICLR 2025. arXiv:2410.02644. https://arxiv.org/abs/2410.02644"),
    ("a_harms", "Chan, A., Salganik, R., Markelius, A., et al. Harms from Increasingly Agentic Algorithmic Systems. FAccT 2023. arXiv:2302.10329. https://arxiv.org/abs/2302.10329"),
    ("a_darkpatterns", "Tang, J., Chen, C., Li, J., et al. Dark Patterns Meet GUI Agents: LLM Agent Susceptibility to Manipulative Interfaces and the Role of Human Oversight. 2025. arXiv:2509.10723. https://arxiv.org/abs/2509.10723"),
    ("a_humancentered", "Chen, C., Zhang, Z., Khalilov, I., et al. Toward a Human-Centered Evaluation Framework for Trustworthy LLM-Powered GUI Agents. 2025. arXiv:2504.17934. https://arxiv.org/abs/2504.17934"),
    ("a_monitoring", "Kale, N., Zhang, C. B. C., Zhu, K., et al. Reliable Weak-to-Strong Monitoring of LLM Agents. 2025. arXiv:2508.19461. https://arxiv.org/abs/2508.19461"),
    ("a_masfail", "Cemri, M., Pan, M. Z., Yang, S., et al. Why Do Multi-Agent LLM Systems Fail? 2025. arXiv:2503.13657. https://arxiv.org/abs/2503.13657"),
    ("a_liedoctor", "Sun, Y., Chen, C., Zhou, Z., et al. It Lied to a Doctor to Buy Poison Ingredients: Quantifying Real-World Misuse of Phone-use Agents. 2026. arXiv:2606.27944. https://arxiv.org/abs/2606.27944"),
    ("a_safeorincapable", "Tang, Z., Zhang, Y., Li, C., et al. Safe, or Simply Incapable? Rethinking Safety Evaluation for Phone-Use Agents. 2026. arXiv:2605.07630. https://arxiv.org/abs/2605.07630"),
    ("a_alignmentlocal", "An, H., Song, Y., Bai, Z., et al. Alignment Is Local: A Paired Diagnostic for GUI Agents under User-Side Persuasion. 2026. arXiv:2607.29199. https://arxiv.org/abs/2607.29199"),
    ("a_safetydrift", "Yu, S., Carroll, F., Bentley, B. L. Operational Hallucination and Safety Drift in AI Agents. IEEE ICAD 2026. arXiv:2607.18366. https://arxiv.org/abs/2607.18366"),
    ("a_oversightcapacity", "Turan, E. Oversight Has a Capacity: Calibrating Agent Guards to a Subjective, Fatiguing Human. 2026. arXiv:2606.08919. https://arxiv.org/abs/2606.08919"),
    ("a_knowingnotenough", "Fu, X., Ramasubbu, N., Galletta, D. Knowing Is Not Enough: Information Retrievability as a Precondition to Effective LLM Oversight. 2026. arXiv:2609.01976. https://arxiv.org/abs/2609.01976"),

    # —— 二、学术文献：无障碍与移动可访问性 ——
    ("a_screenrecog", "Zhang, X., de Greef, L., Swearngin, A., et al. Screen Recognition: Creating Accessibility Metadata for Mobile Applications from Pixels. 2021. arXiv:2101.04893. https://arxiv.org/abs/2101.04893"),
    ("a_labeldroid", "Chen, J., Chen, C., Xing, Z., et al. Unblind Your Apps: Predicting Natural-Language Labels for Mobile GUI Components by Deep Learning. ICSE 2020. arXiv:2003.00380. https://arxiv.org/abs/2003.00380"),
    ("a_a11yinvestigation", "Chen, S., Chen, C., Fan, L., et al. Accessible or Not? An Empirical Investigation of Android App Accessibility. 2022. arXiv:2203.06422. https://arxiv.org/abs/2203.06422"),
    ("a_alticon", "Haque, S., Csallner, C. Early Accessibility: Automating Alt-Text Generation for UI Icons During App Development. 2025. arXiv:2504.13069. https://arxiv.org/abs/2504.13069"),
    ("a_insight", "Ansah, J. O., Kapoor, A., Khanna, A., et al. Insight: Enhancing Mobile Accessibility for Blind and Visually Impaired Users with LLMs. 2026. arXiv:2605.09803. https://arxiv.org/abs/2605.09803"),
    ("a_notana11y", "Deivasigamani, R., Alvi, S. F., Andrea, D., et al. Not an A11y: How Android Accessibility Exposes Mobile AI Agents to Indirect Prompt Injection. 2026. arXiv:2608.08939. https://arxiv.org/abs/2608.08939"),
    ("a_virtualization", "Minn, W., Phan, P., Malviya, V. K., et al. Virtualization-based Penetration Testing Study for Detecting Accessibility Abuse Vulnerabilities in Banking Apps in East and Southeast Asia. APSEC 2025 SEIP. arXiv:2601.21258. https://arxiv.org/abs/2601.21258"),
    ("a_mcp_a11y", "Ramineni, V., Saksena, N., Agarwal, A. K., et al. MCP-Driven Accessibility Tree Standardization for AI-Powered Screen Reader Agents. 2026. arXiv:2608.24898. https://arxiv.org/abs/2608.24898"),

    # —— 二、学术文献：面向老年人的 AI 与数字鸿沟（2025—2026） ——
    ("a_elderbench", "Zhan, W., Shaqu, Q., Liu, Y., et al. ElderBench: Benchmarking Autonomous Mobile Agents for Older Adults. 复旦大学. 2026-09-04. arXiv:2609.04850（arXiv 预印本，19 页，无会议信息；注意与同名 IEEE 工作“ElderBench: Benchmarking Personalized Open-Source LLMs for Older Adults”非同一篇）. https://arxiv.org/abs/2609.04850"),
    ("a_divideusability", "Ine, C. The Digital Divide in Geriatric Care: Why Usability, Not Access, is the Real Problem. 2026. arXiv:2601.17012. https://arxiv.org/abs/2601.17012"),
    ("a_helpinghelper", "Sharifi, H., Shomee, H. H., Lamar, M., et al. Helping the Helper: LLM-Assisted Problem Articulation for Older Adults Seeking Technology Support. 2026. arXiv:2601.10018. https://arxiv.org/abs/2601.10018"),
    ("a_cogbridging", "Chen, F., Li, L. X., Chung, R.-Y., et al. Bridging the Cognitive Gap: Co-Designing and Evaluating a Voice-Enabled Community Chatbot for Older Adults. 2026. arXiv:2603.11303. https://arxiv.org/abs/2603.11303"),
    ("a_agreeableness", "Mathur, N., Rahman, H., Desai, S. “Who wants to be nagged by AI?”: Investigating the Effects of Agreeableness on Older Adults’ Perception of LLM-Based Voice Assistants’ Explanations. CHI 2026. arXiv:2603.09012. https://arxiv.org/abs/2603.09012"),
    ("a_convbreakdown", "Mathur, N., Zubatiy, T., Rozga, A., et al. “It feels like hard work trying to talk to it”: Understanding Older Adults’ Experiences of Encountering and Repairing Conversational Breakdowns with AI Systems. 2025. arXiv:2510.06690. https://arxiv.org/abs/2510.06690"),
    ("a_eldercareagentic", "Khalil, R. A., Ahmad, K., Ali, H. Redefining Elderly Care with Agentic AI: Challenges and Opportunities. 2025. arXiv:2507.14912. https://arxiv.org/abs/2507.14912"),
    ("a_wepilot", "Zhang, H., Zhang, P., Chen, Y., et al. WePilot: Integrating Younger Family Members and Chatbot to Support Older Adults Learning Smartphone Usage. Proc. ACM Hum.-Comput. Interact. 9(7), CSCW522, 2025. DOI:10.1145/3757703. https://doi.org/10.1145/3757703"),
    ("a_agemate", "Chen, L., Mu, Y. Bridging the Digital Divide: Empowering Elderly Smartphone Users with Intelligent and Human-Centered Design in Agemate. Proceedings of the 1st Workshop for Research on Agent Language Models (REALM 2025). DOI:10.18653/v1/2025.realm-1.23. https://aclanthology.org/2025.realm-1.23/"),
    ("a_exploar", "Li, J., Qiu, L., Wu, Z., et al. ExplorAR: Assisting Older Adults to Learn Smartphone Apps through AR-powered Trial-and-Error with Interactive Guidance. ACM Multimedia 2025. DOI:10.1145/3746027.3755578. arXiv:2508.01282. https://arxiv.org/abs/2508.01282"),
    ("a_olla", "Kodandaram, S. R., Reddy, M. P., Bi, X., et al. Are We There Yet? Assessing Computer-Use Agents for Blind Users’ Accessible Interaction with Desktop Applications. EMNLP 2026. arXiv:2609.00524. https://arxiv.org/abs/2609.00524"),
    ("a_autobalance", "Karimi, P., Martin-Hammond, A. Finding the Right Balance: User Control and Automation in AI Tools for Supporting Older Adults’ Health Information Tasks. CHI Extended Abstracts 2025. DOI:10.1145/3706599.3719773. https://doi.org/10.1145/3706599.3719773"),
    ("a_autoboundary", "You, L., Zhou, J. Designing Automation Boundaries for Trustworthy Smart Medication Support. 2026. arXiv:2606.28777. https://arxiv.org/abs/2606.28777"),
    ("a_oversightembed", "Meng, Y., Chen, J., Ye, L., et al. Balancing Safety and Autonomy: Accessibility-Oriented Interventions in Generative AI for Cognitive Impairment. ASSETS 2026. DOI:10.1145/3797867.3829017. arXiv:2608.17175. https://arxiv.org/abs/2608.17175"),
    ("a_intergen", "Chu, H., Lee, Y., Park, Y. S., et al. When AI “Works,” When Does Help Begin?: Intergenerational Support Around Older Adults’ LLM Usage. CSCW 2026 workshop. arXiv:2608.24297. https://arxiv.org/abs/2608.24297"),
    ("a_grandguard", "Fan, C., Yang, X., Zheng, Y., et al. GrandGuard: Taxonomy, Benchmark, and Safeguards for Elderly-Chatbot Interaction Safety. Findings of ACL 2026. DOI:10.18653/v1/2026.findings-acl.1116. arXiv:2605.20203. https://arxiv.org/abs/2605.20203"),
    ("a_nexui", "Patapati, S. Navigation Alone Is Not Enough: Evaluating Explanatory Assistive UI Agents. 2026. arXiv:2608.09944. https://arxiv.org/abs/2608.09944"),

    # —— 三、技术文档 ——
    ("t_android", "Android Developers. AccessibilityService / AccessibilityService.takeScreenshot / 包可见性（<queries>）/ 前台服务类型. https://developer.android.com/reference/android/accessibilityservice/AccessibilityService"),
    ("t_openai", "OpenAI. Chat Completions API 与 function / tool calling 文档. https://platform.openai.com/docs/api-reference/chat"),
    ("t_deepseek", "DeepSeek. API 文档（上下文硬盘缓存）. https://api-docs.deepseek.com/"),
    ("t_xfyun", "科大讯飞开放平台. 语音听写（流式版）WebSocket API 文档. https://www.xfyun.cn/doc/asr/voicedictation/API.html"),
]

REF_GROUPS = [
    ("一、政策与统计文献", "p_"),
    ("二、学术文献（GUI 智能体及相关研究）", "a_"),
    ("三、技术文档", "t_"),
]

CITE_RE = re.compile(r"\[@([a-z0-9_]+)\]")


# 列宽（cm，合计 15cm = A4 正文宽度）。仅对列数多、首列文字长的表做定制，
# 否则 7 列表格每列仅约 2cm，方案名会被挤成多行。
# 列宽按**题注名**索引，不按表号——表号会因插入新表而整体后移，
# 用表号做键曾导致多处定制列宽静默失效（退回等宽）。题注名是稳定的。
# 合计 15cm = A4 正文宽度；仅对列数多、或某列文字明显更长的表做定制。
COL_WIDTHS = {
    "老年人手机办事痛点的四层十二项剖析": [1.6, 1.8, 3.4, 2.6, 2.6, 3.0],
    "现有助老方案对比（一）：执行能力维度": [4.2, 1.8, 1.8, 1.8, 1.8, 1.8, 1.8],
    "现有助老方案对比（二）：工程与合规维度": [4.2, 1.8, 1.8, 1.8, 1.8, 1.8, 1.8],
    "面向老年人的需求清单与功能对应": [1.4, 3.4, 2.8, 3.4, 1.8, 2.2],
    "已有工作与本作品的技术差异": [2.4, 4.6, 4.6, 3.4],
    "八项创新点及其差异与支撑证据": [3.0, 2.6, 4.6, 2.8, 2.0],
    "典型任务全流程走查（美团点餐至结算交接）": [1.5, 3.3, 3.4, 4.8, 2.0],
    "跨应用悬浮接线台的交互要素与工程细节": [3.6, 11.4],
}


def build_ref_numbers():
    """按 REFS 顺序分配编号；返回 {key: 序号}。"""
    return {key: i + 1 for i, (key, _) in enumerate(REFS)}


# ---------------------------------------------------------------- 基础工具

def set_font(run, name=FONT, size=None, bold=None, color=None, latin=None):
    """按中西文分槽设置字体。

    中文写进 w:eastAsia，西文/数字/半角标点写进 w:ascii 与 w:hAnsi。
    代码之类的等宽场景传入 latin=name，四个槽统一为等宽字体，保持对齐。
    """
    latin = latin if latin is not None else (name if name == MONO else LATIN)
    run.font.name = latin
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(a), latin)


def spacing(p, before=0, after=0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return p


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+?`)")


def add_runs(p, text, size=BODY, bold=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            set_font(p.add_run(part[2:-2]), FONT, size, True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            set_font(p.add_run(part[1:-1]), MONO, size * 0.94, bold)
        else:
            set_font(p.add_run(part), FONT, size, bold)
    return p


def add_field(paragraph, instr, placeholder=""):
    run = paragraph.add_run()
    set_font(run, FONT, BODY)
    for tag, attr, text in (("begin", None, None), ("instr", None, instr),
                            ("separate", None, None), ("text", None, placeholder),
                            ("end", None, None)):
        if tag == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        elif tag == "text":
            el = OxmlElement("w:t")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        run._r.append(el)


# ---------------------------------------------------------------- 文档骨架

def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(3)
    sec.header_distance = sec.footer_distance = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY)
    _nf = normal.element.rPr.rFonts
    for _a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if _nf.get(qn(_a)) is not None:
            del _nf.attrib[qn(_a)]
    _nf.set(qn("w:eastAsia"), FONT)
    for _a in ("w:ascii", "w:hAnsi", "w:cs"):
        _nf.set(qn(_a), LATIN)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)

    # 页眉样式默认自带「居中 4680 / 右 9360」两个制表位；段落里再加右制表位时，
    # 制表符会先落到样式继承来的居中位，文字便停在页中。这里直接改写页眉样式：
    # 只保留右制表位（模板 w:tab pos=8306 twips），并清除原有的两个。
    _hs = doc.styles["Header"]
    _hpPr = _hs.element.get_or_add_pPr()
    _htabs = _hpPr.find(qn("w:tabs"))
    if _htabs is not None:
        _hpPr.remove(_htabs)
    _htabs = OxmlElement("w:tabs")
    for _pos in ("4680", "9360"):
        _c = OxmlElement("w:tab")
        _c.set(qn("w:val"), "clear")
        _c.set(qn("w:pos"), _pos)
        _htabs.append(_c)
    _r = OxmlElement("w:tab")
    _r.set(qn("w:val"), "right")
    _r.set(qn("w:pos"), "8306")
    _r.set(qn("w:leader"), "none")
    _htabs.append(_r)
    _hpPr.insert_element_before(
        _htabs, "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
        "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
        "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
        "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
        "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
        "w:pPrChange")

    for name, size in (("Heading 1", H1), ("Heading 2", H2), ("Heading 3", H3)):
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        _sf = st.element.rPr.rFonts
        for _a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if _sf.get(qn(_a)) is not None:
                del _sf.attrib[qn(_a)]
        _sf.set(qn("w:eastAsia"), FONT)
        for _a in ("w:ascii", "w:hAnsi", "w:cs"):
            _sf.set(qn(_a), LATIN)
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        st.paragraph_format.space_after = Pt(6 if name == "Heading 1" else 4)
        st.paragraph_format.keep_with_next = True

    # 页眉照搬模板：大赛 Logo + 「2026第八届全球校园人工智能算法精英大赛」，五号、左对齐、下方横线
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if LOGO.exists():
        hp.add_run().add_picture(str(LOGO), width=Cm(1.117))   # 模板 wp:extent 401955 EMU = 1.117 cm
    # 制表位取自模板：居中 4153、右对齐 8306 twips（正文宽 15cm）
    # 模板定义了居中(4153)与右对齐(8306)两个制表位；若两者并存，单个 <w:tab/> 会先落到居中位，
    # 文字就停在页中。模板自身是靠 32 个空格把文字推右的，且会越过页眉横线压进右边距。
    # 这里只保留模板定义的右制表位，既实现“文字靠右”，又不越界。
    _tabs = OxmlElement("w:tabs")
    _t = OxmlElement("w:tab")
    _t.set(qn("w:val"), "right")
    _t.set(qn("w:pos"), "8306")
    _t.set(qn("w:leader"), "none")
    _tabs.append(_t)
    hp._p.get_or_add_pPr().insert_element_before(
        _tabs, "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
        "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
        "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
        "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
        "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
        "w:pPrChange")
    hp.add_run().add_tab()
    set_font(hp.add_run(HEADER_TEXT), FONT, S["五号"], latin=LATIN)
    _bdr = OxmlElement("w:pBdr")
    _bot = OxmlElement("w:bottom")
    for _k, _v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", "000000")):
        _bot.set(qn(_k), _v)
    _bdr.append(_bot)
    # pBdr 必须排在 spacing / ind / jc 之前（OOXML 架构顺序）
    hp._p.get_or_add_pPr().insert_element_before(
        _bdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
        "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
        "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
        "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
        "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
        "w:pPrChange",
    )

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE   \\* MERGEFORMAT")

    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    doc.settings.element.append(uf)
    return doc


def build_cover(doc, title, subtitle, infos):
    """封面严格照搬大赛模板的排版。

    规格取自模板样式表（用 LibreOffice 转换后逐段读取）：
      标题区 黑体，22/20/18pt，居中；信息区 仿宋 16pt，首行缩进 7.25 字符。
    """
    def line(text, size, font, align=None, indent_chars=None, before=0, after=0):
        p = spacing(doc.add_paragraph(), before, after)
        # 模板封面继承 Normal 的 w:line="360"（1.5 倍），照搬其观感
        p.paragraph_format.line_spacing = 1.5
        if align is not None:
            p.alignment = align
        if indent_chars is not None:
            ind = p._p.get_or_add_pPr().get_or_add_ind()
            ind.set(qn("w:firstLineChars"), str(indent_chars))
            ind.set(qn("w:firstLine"), "0")
        if text:
            # 封面西文与数字同用中文字体，与模板一致
            set_font(p.add_run(text), font, size, latin=font)
        else:
            # 空行的行高由“段落标记”的字符格式决定；只加空 run 不起作用，
            # 必须把字号写进 pPr/rPr，否则空行会比模板矮一截。
            _pPr = p._p.get_or_add_pPr()
            rPr = _pPr.find(qn("w:rPr"))
            if rPr is None:
                rPr = OxmlElement("w:rPr")
                # pPr 内 rPr 排在最后（仅 sectPr / pPrChange 之后）
                _pPr.insert_element_before(rPr, "w:sectPr", "w:pPrChange")
            rf = rPr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rPr.insert(0, rf)
            rf.set(qn("w:eastAsia"), font)
            for _a in ("w:ascii", "w:hAnsi", "w:cs"):
                rf.set(qn(_a), font)
            for _tag in ("w:sz", "w:szCs"):
                _e = OxmlElement(_tag)
                _e.set(qn("w:val"), str(int(size * 2)))
                rPr.append(_e)
        return p

    C, L = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT

    def info_line(text, underline):
        """信息行：标签不划线、填写内容划线；整段左缩进 7.25 字符，换行后仍对齐。"""
        p = spacing(doc.add_paragraph(), 0, 0)
        p.paragraph_format.line_spacing = 1.5
        p.alignment = L
        ind = p._p.get_or_add_pPr().get_or_add_ind()
        ind.set(qn("w:leftChars"), "725")     # 用左缩进而非首行缩进，长内容换行不会跑到页边
        ind.set(qn("w:firstLineChars"), "0")
        ind.set(qn("w:firstLine"), "0")
        label, _, value = text.partition("：")
        set_font(p.add_run(label + "：" if value else label), FANG, 16, latin=FANG)
        if value:
            r = p.add_run(value)
            set_font(r, FANG, 16, latin=FANG)
            if underline:
                r.font.underline = True
        elif underline:
            r = p.add_run(" " * 13)
            set_font(r, FANG, 16, latin=FANG)
            r.font.underline = True
        return p

    # 空行高度按段落标记字号计算。模板用 3 个空行，这里减为 2 个，
    # 是为了把下方空间腾出来让“日期”留在同一页。
    for _ in range(2):
        line("", 22, HEI, align=C)
    line("2026年第八届", 22, HEI, align=C, indent_chars=0)
    line("全球校园人工智能算法精英大赛", 22, HEI, align=C, indent_chars=0)
    line("", 22, HEI)
    line("算法创新赛", 20, HEI, align=C, indent_chars=0)
    # 赛题名称：沿用 Markdown 的副标题，去掉其中的“技术报告”字样
    topic = re.sub(r"\s*技\s*术\s*报\s*告\s*$", "", subtitle).strip()
    line(topic, 18, HEI, align=C, indent_chars=0)
    line("技术报告", 20, HEI, align=C, indent_chars=0)
    line("", 22, HEI)
    # 前三条信息带下划线，其后再空 6 行，最后是日期（模板中日期不划线）
    for info in infos[:3]:
        info_line(info, underline=True)
    # 模板在此处留 6 个空行，但 1.5 倍行距下会把日期挤到第 2 页；
    # 减为 2 个，保证封面（含日期）完整落在第一页内。
    for _ in range(2):
        line("", 22, HEI)
    if len(infos) >= 4:
        info_line(infos[3], underline=False)


# ---------------------------------------------------------------- 正文渲染

def add_table(doc, rows, caption, number):
    cap = spacing(doc.add_paragraph(), 6, 2)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    add_runs(cap, f"表 {number}　{caption}", NOTE, True)

    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = COL_WIDTHS.get(caption)
    if widths and len(widths) == len(header):
        t.autofit = False
    else:
        t.autofit = True

    for i, text in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = spacing(cell.paragraphs[0], 2, 2)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, text, TBL, True)
        shade_cell(cell, "F2F2F2")

    for row in body:
        cells = t.add_row().cells
        for i, text in enumerate(row[: len(header)]):
            cell = cells[i]
            cell.text = ""
            first = True
            for seg in text.split("<br>"):
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                spacing(p, 1, 1)
                add_runs(p, seg, TBL)

    if widths and len(widths) == len(header):
        # 必须同时写 tblGrid（列宽）与 tcW（单元格宽）：python-docx 的 cell.width
        # 只写 tcW，而 Word / WPS / LibreOffice 都优先按 tblGrid 排版，
        # 不写 tblGrid 的话定制列宽不会生效（曾因此静默退回等宽）。
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        t._tbl.tblPr.append(lay)

    spacing(doc.add_paragraph(), 0, 4).add_run("")


MAX_FIG_W_CM = 15.0   # A4 正文宽度
MAX_FIG_H_CM = 19.0   # 留出题注与页边距；过高会让图上的字印出来过小


def add_figure(doc, png, caption, number, max_h_cm=MAX_FIG_H_CM):
    """按宽高双约束缩放后插入。

    手机截图是竖版（如 1272×2800），只按宽度设为 15cm 会得到约 33cm 的高度、
    直接超出一页正文，排版会散架。因此先按宽度算，若算出的高度超限就改按高度反推宽度。
    """
    try:
        from PIL import Image
        with Image.open(png) as im:
            w_px, h_px = im.size
    except Exception:
        w_px = h_px = 0

    width_cm = MAX_FIG_W_CM
    if w_px and h_px:
        height_cm = width_cm * h_px / w_px
        if height_cm > max_h_cm:
            width_cm = max_h_cm * w_px / h_px

    p = spacing(doc.add_paragraph(), 6, 2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(png), width=Cm(width_cm))
    cap = spacing(doc.add_paragraph(), 2, 10)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(cap, f"图 {number}　{caption}", NOTE, True)


def render(doc, text, num):
    lines = text.split("\n")
    i = 0
    pending = []
    tno = fno = 0          # fno：Mermaid 图序号（对应 figN.png 与 FIGURE_CAPTIONS）
    shot = 0               # 真机截图数量
    fdisp = 0              # “图 N”的显示序号，按文档出现顺序
    cited = set()

    def expand(s):
        """把 [@key] 换成 [序号]，并记录引用过哪些文献。"""
        def sub(m):
            key = m.group(1)
            if key not in num:
                raise KeyError(f"正文引用了未定义的文献键：{key}")
            cited.add(key)
            return f"[{num[key]}]"
        return CITE_RE.sub(sub, s)

    def flush():
        nonlocal pending, tno
        if pending:
            tno += 1
            cap = TABLE_CAPTIONS[tno - 1] if tno <= len(TABLE_CAPTIONS) else "（缺题注）"
            add_table(doc, pending, cap, tno)
            pending = []

    while i < len(lines):
        raw = lines[i]
        s = raw.strip()

        if s.startswith("|") and s.endswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                pending.append([expand(c) for c in cells])
            i += 1
            continue
        flush()

        if not s or s == "---" or s.startswith("<!--"):
            i += 1
            continue

        # 参考文献表：按分组顺序自动生成，编号与正文一致
        if s == "[REFERENCES]":
            for group_title, prefix in REF_GROUPS:
                items = [(k, t) for k, t in REFS if k.startswith(prefix)]
                if not items:
                    continue
                gp = spacing(doc.add_paragraph(), 8, 3)
                gp.paragraph_format.keep_with_next = True
                set_font(gp.add_run(group_title), FONT, S["小四"], True)
                for k, t in items:
                    rp = spacing(doc.add_paragraph(), 0, 2)
                    rp.paragraph_format.left_indent = Cm(0.9)
                    rp.paragraph_format.first_line_indent = Cm(-0.9)
                    add_runs(rp, f"[{num[k]}] {t}", S["五号"])
            i += 1
            continue

        # 真机截图：![题注](相对路径)。题注取 alt 文字，编号沿用同一个图表计数器，
        # 与 Mermaid 图共用"图 N"序列。
        mi = re.match(r"^!\[(.*?)\]\((.+?)\)\s*$", s)
        if mi:
            alt, rel = mi.group(1), mi.group(2)
            png = (HERE / rel)
            shot += 1
            fdisp += 1
            if png.exists():
                add_figure(doc, png, alt, fdisp)
            else:
                pp = spacing(doc.add_paragraph(), 6, 10)
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(pp, f"图 {fdisp}　（缺少 {rel}）", NOTE)
            i += 1
            continue

        # Mermaid 图：插入已渲染的 PNG，题注在图下方
        if s.startswith("```mermaid"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                i += 1
            i += 1
            fno += 1
            fdisp += 1
            png = FIGDIR / f"fig{fno}.png"
            if png.exists():
                cap = FIGURE_CAPTIONS[fno - 1] if fno <= len(FIGURE_CAPTIONS) else ""
                add_figure(doc, png, cap, fdisp)
            else:
                p = spacing(doc.add_paragraph(), 6, 10)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(p, f"图 {fdisp}　（缺少 {png.name}，请先运行 render_figures.sh）", NOTE)
            continue

        # 其他代码块
        if s.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                p = spacing(doc.add_paragraph(), 0, 0)
                p.paragraph_format.left_indent = Cm(0.5)
                set_font(p.add_run(lines[i]), MONO, NOTE)
                i += 1
            i += 1
            spacing(doc.add_paragraph(), 0, 4).add_run("")
            continue

        # 引用块（边界声明等）
        if s.startswith(">"):
            p = spacing(doc.add_paragraph(), 4, 4)
            p.paragraph_format.left_indent = Cm(0.75)
            add_runs(p, "▍" + expand(s.lstrip("> ").strip()), BODY)
            i += 1
            continue

        # 标题
        for marker, lvl, size in (("#### ", 3, H3), ("### ", 2, H2), ("## ", 1, H1)):
            if s.startswith(marker):
                p = doc.add_heading(level=lvl)
                add_runs(p, expand(s[len(marker):]), size, True)
                spacing(p, 12 if lvl == 1 else 8, 6 if lvl == 1 else 4)
                break
        else:
            # 列表
            if s.startswith("- "):
                p = spacing(doc.add_paragraph(), 0, 2)
                p.paragraph_format.left_indent = Cm(0.74)
                p.paragraph_format.first_line_indent = Cm(-0.37)
                add_runs(p, "· " + expand(s[2:]), BODY)
            elif re.match(r"^\d+\.\s", s):
                p = spacing(doc.add_paragraph(), 0, 2)
                p.paragraph_format.left_indent = Cm(0.74)
                p.paragraph_format.first_line_indent = Cm(-0.37)
                add_runs(p, expand(s), BODY)
            else:
                p = spacing(doc.add_paragraph(), 0, 3)
                p.paragraph_format.first_line_indent = Pt(BODY * 2)
                add_runs(p, expand(s), BODY)
        i += 1

    flush()
    return tno, fdisp, cited, shot


def main():
    src = SRC.read_text(encoding="utf-8")
    src = re.sub(r"<!--.*?-->", "", src, flags=re.S)
    lines = src.split("\n")

    title = subtitle = ""
    infos, body = [], []
    in_cover = False           # 只有封面区（副标题之后、第一条 --- 之前）的 - 行才是封面信息
    for line in lines:
        s = line.strip()
        if s.startswith("# ") and not title:
            title = s[2:].strip()
            continue
        if s.startswith("## ") and title and not subtitle:
            subtitle = s[3:].strip()
            in_cover = True
            continue
        if in_cover:
            if s == "---":
                in_cover = False
                body.append(line)
            elif s.startswith("- "):
                infos.append(s[2:].strip())
            else:
                body.append(line)
            continue
        body.append(line)

    if len(infos) != 4:
        print(f"!! 封面信息行数为 {len(infos)}，预期 4（团队名称/参赛编号/作品名称/日期）")
        for x in infos[:10]:
            print("   ", x[:60])
        return 1

    doc = new_document()
    build_cover(doc, title, subtitle, infos)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    p = spacing(doc.add_paragraph(), 0, 12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(p, "目　录", S["二号"], True)
    add_field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u',
              "（在 Word / WPS 中按 F9 或右键“更新域”生成目录）")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    num = build_ref_numbers()
    ntab, nfig, cited, nshot = render(doc, "\n".join(body), num)
    doc.save(OUT)

    print(f"输出：{OUT}")
    print(f"表格 {ntab} 张（题注 {len(TABLE_CAPTIONS)} 条）｜插图 {nfig} 张（其中 Mermaid 流程图 {len(FIGURE_CAPTIONS)} 张、图片 {nshot} 张）")
    print(f"参考文献 {len(REFS)} 条，正文引用 {len(cited)} 条")
    unused = [k for k, _ in REFS if k not in cited]
    if unused:
        print(f"!! 以下文献从未在正文引用（建议删除或补引用）：{unused}")
    if ntab != len(TABLE_CAPTIONS):
        print(f"!! 表格数量与题注数量不一致：{ntab} vs {len(TABLE_CAPTIONS)}")
        return 1
    if nfig != len(FIGURE_CAPTIONS) + nshot:
        print(f"!! 插图数量与题注数量不一致：{nfig} vs {len(FIGURE_CAPTIONS)}+{nshot}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
